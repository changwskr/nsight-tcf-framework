from pathlib import Path
import hashlib
import re
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

A=Path('C:/Programming(23-08-15)/nsight-tcf-framework/ztcf-다이어리/2026-08-17-알아야 되는 아키텍처')
B=Path('C:/Programming(23-08-15)/nsight-tcf-framework/pdmg-service/docs')
names=['네이밍원칙.md','book/source/네이밍원칙.md','book/source/53.네이밍 아키텍처 다이어그램.md','book/source/06.네이밍 형식.md','book/source/06.네이밍 형식-1.md','book/chapter/05장.네이밍 규칙.md','book/chapter-확장본/05장.네이밍_규칙_ASCII_확장본.md','53.네이밍 아키텍처 다이어그램.md','06.네이밍 형식.md','06.네이밍 형식-1.md']
paths=[A/'pdg-fw-docs/Java 네이밍 구조 + 주석 처리  표준.md']+[A/'pdmg-service-docs'/n for n in names]+[B/n for n in names]
groups={}
for p in paths:
    raw=p.read_bytes()
    digest=hashlib.sha256(raw).hexdigest()
    groups.setdefault(digest, {'paths':[], 'text':raw.decode('utf-8-sig')})['paths'].append(p)
d=Document()
s=d.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet']:
    st=d.styles[name]
    st.font.name='Calibri'
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
d.styles['Normal'].font.size=Pt(11)
for name,size in [('Heading 1',16),('Heading 2',13),('Heading 3',12)]:
    d.styles[name].font.size=Pt(size)
d.add_heading('Java 네이밍 및 주석 통합 표준서',0)
d.add_paragraph('제공 문서 전체 수록 · 원문 대조용 통합본', 'Subtitle')
d.add_paragraph('작성일: 2026-08-28')
d.add_heading('편집 범위와 읽는 방법',1)
d.add_paragraph(f'지정된 {len(paths)}개 파일을 모두 읽고 SHA-256 기준으로 동일한 내용을 묶었습니다. 서로 다른 {len(groups)}개 원문 버전은 각각 독립된 장으로 수록하였습니다. 경로가 다른 동일 문서는 중복 수록하지 않으며, 내용이 조금이라도 다른 문서는 유지합니다.')
d.add_paragraph('이 문서는 제공된 네이밍·주석 자료의 통합 열람본입니다. 문서 사이의 상충하는 규칙을 임의로 확정하거나 최신 표준으로 변경하지 않았습니다. 실제 적용 시 프로젝트에서 승인한 기준을 확인하십시오. 코드와 텍스트 다이어그램은 원문 표기를 유지합니다.')
d.add_heading('핵심 적용 관점',1)
for t in ['업무 코드와 서비스 식별자의 규칙을 Java 클래스·메서드·변수 명명 규칙과 구분합니다.', '패키지, Controller, Service, DAO, DTO의 명명 규칙은 동일한 업무 식별 체계를 바탕으로 대조합니다.', '동일 제목의 문서라도 경로별 버전 차이가 있으므로 아래 원문 목록과 함께 검토합니다.', '주석 및 예시 코드는 원문 자료이며, 이 통합본 자체가 신규 구현이나 규칙 변경을 승인하지는 않습니다.']:
    d.add_paragraph(t,'List Bullet')
d.add_heading('원문 파일 목록',1)
for i,(digest,g) in enumerate(groups.items(),1):
    d.add_heading(f'원문 {i:02d} · {g["paths"][0].name}',2)
    d.add_paragraph('SHA-256: '+digest)
    for p in g['paths']: d.add_paragraph(str(p),'List Bullet')

def code_line(line):
    p=d.add_paragraph()
    p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.line_spacing=1
    r=p.add_run(line)
    r.font.name='Consolas'; r.font.size=Pt(8)
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')

line_count=0
for i,(digest,g) in enumerate(groups.items(),1):
    d.add_page_break()
    d.add_heading(f'원문 {i:02d} · {g["paths"][0].name}',1)
    d.add_paragraph('대표 출처: '+str(g['paths'][0]))
    d.add_paragraph('아래는 해당 버전의 전체 내용입니다. Markdown 표와 코드 블록은 원문 텍스트로 보존합니다.')
    fenced=False
    for line in g['text'].splitlines():
        line_count+=1
        if line.lstrip().startswith('```') or line.lstrip().startswith('~~~'):
            fenced=not fenced
            code_line(line)
        elif fenced or line.lstrip().startswith('|'):
            code_line(line)
        elif re.match(r'^#{1,6} ',line):
            n=len(line)-len(line.lstrip('#'))
            d.add_heading(line[n:].strip(),min(n+1,3))
        elif re.match(r'^\s*[-*+] ',line):
            d.add_paragraph(line,'List Bullet')
        else:
            d.add_paragraph(line)
footer=s.footer.paragraphs[0]
footer.alignment=2
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
out=A/'Java_네이밍_및_주석_통합_표준서.docx'
d.save(out)
# Verify the document package and complete line-by-line source inclusion.
from zipfile import ZipFile
from lxml import etree
with ZipFile(out) as z:
    assert z.testzip() is None
    root=etree.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    paragraphs=[''.join(p.itertext()) for p in []]
    paragraphs=[''.join(p.xpath('.//w:t/text()',namespaces=ns)) for p in root.xpath('//w:p',namespaces=ns)]
    expected=[]
    for g in groups.values():
        for line in g['text'].splitlines():
            if re.match(r'^#{1,6} ',line): line=line.lstrip('#').strip()
            expected.append(line)
    # Each nonempty source line is retained; repeated lines are counted as well.
    from collections import Counter
    actual=Counter(paragraphs); required=Counter(expected)
    missing={k:v-actual[k] for k,v in required.items() if k and actual[k]<v}
    assert not missing, str(missing)[:500]
print(f'OUTPUT: {out}\nSOURCES: {len(paths)}\nUNIQUE: {len(groups)}\nSOURCE_LINES: {line_count}\nBYTES: {out.stat().st_size}\nZIP_AND_TEXT_CHECK: PASS')
