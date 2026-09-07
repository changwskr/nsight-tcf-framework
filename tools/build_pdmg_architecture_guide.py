from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(r"C:\Programming(23-08-15)\nsight-tcf-framework\ztcf-다이어리\2026-09-01-PDMG 아키텍처 정의서\아키텍처정의서(T5-Final)\NSIGHT_PDMG_ARCHITECTURE_BOOK_FINAL_V5_COMPLETE\PDMG_아키텍처_수립_가이드.docx")

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
PALE_BLUE = "E8EEF5"
LIGHT = "F4F6F9"
GRAY = "666666"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
GOLD = "B38B2E"
RED = "9B1C1C"
GREEN = "2F6B4F"
FONT = "Malgun Gothic"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    total = sum(widths_dxa)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_para(p, before=0, after=6, line=1.25, keep=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep


def add_p(doc, text="", bold_prefix=None, italic=False, color=None, after=6, align=None):
    p = doc.add_paragraph()
    style_para(p, after=after)
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True, color=color)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, color=color, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, color=color, italic=italic)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    style_para(p, after=4, line=1.25)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    style_para(p, after=4, line=1.25)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    for r in p.runs:
        set_run_font(r)
    return p


def add_callout(doc, label, text, kind="info"):
    fill = {"info": PALE_BLUE, "risk": "FCE8E6", "success": "E8F2EC", "note": LIGHT}.get(kind, LIGHT)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    style_para(p, after=0, line=1.2)
    r = p.add_run(f"{label}  ")
    set_run_font(r, bold=True, color=NAVY if kind != "risk" else RED)
    r = p.add_run(text)
    set_run_font(r)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_table(doc, headers, rows, widths, header_fill=PALE_BLUE, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_para(p, after=0, line=1.05)
        r = p.add_run(h)
        set_run_font(r, size=font_size, bold=True, color=NAVY)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            p = cells[i].paragraphs[0]
            style_para(p, after=0, line=1.1)
            r = p.add_run(str(val))
            set_run_font(r, size=font_size)
            if i == 0 and len(headers) <= 3:
                r.bold = True
        set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, size=9, color=GRAY)


def add_toc_field(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Word에서 문서를 열고 목차를 선택한 후 [필드 업데이트]를 실행하십시오."
    fld3 = OxmlElement("w:fldChar")
    fld3.set(qn("w:fldCharType"), "end")
    run._r.extend([fld1, instr, fld2, text, fld3])
    set_run_font(run, size=10, color=GRAY)


def add_section_intro(doc, number, title, purpose, source_chapters):
    doc.add_page_break()
    p = doc.add_paragraph()
    style_para(p, before=40, after=10)
    r = p.add_run(f"PART {number}")
    set_run_font(r, size=11, bold=True, color=GOLD)
    p = doc.add_paragraph()
    style_para(p, after=12)
    r = p.add_run(title)
    set_run_font(r, size=25, bold=True, color=NAVY)
    add_callout(doc, "이 부의 목적", purpose, "info")
    add_p(doc, f"근거 장: {source_chapters}", italic=True, color=GRAY, after=16)


def add_gate(doc, rows):
    add_heading(doc, "Gate 판정", 3)
    add_table(doc, ["판정", "적용 기준"], rows, [1700, 7660])


def add_phase(doc, idx, name, tagline, inputs, activities, outputs, gate, forbidden, evidence):
    doc.add_page_break()
    add_heading(doc, f"{idx}단계. {name}", 1)
    add_callout(doc, "단계 정의", tagline, "info")
    add_heading(doc, "핵심 질문", 2)
    for x in activities["questions"]:
        add_bullet(doc, x)
    add_heading(doc, "입력", 2)
    add_table(doc, ["입력 항목", "확인 관점"], inputs, [2600, 6760])
    add_heading(doc, "수행 절차", 2)
    for x in activities["steps"]:
        add_number(doc, x)
    add_heading(doc, "핵심 의사결정", 2)
    for x in activities["decisions"]:
        add_bullet(doc, x)
    add_heading(doc, "필수 산출물", 2)
    add_table(doc, ["산출물", "완료 조건"], outputs, [2600, 6760])
    add_gate(doc, [["PASS", gate[0]], ["CONDITIONAL", gate[1]], ["FAIL", gate[2]]])
    add_heading(doc, "정상 패턴과 금지 패턴", 2)
    add_table(doc, ["정상 패턴", "금지 패턴"], [[forbidden[0], forbidden[1]]], [4680, 4680], header_fill=LIGHT_GRAY)
    add_heading(doc, "Evidence 패키지", 2)
    for x in evidence:
        add_bullet(doc, x)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [
        ("Title", 30, NAVY, 0, 8),
        ("Subtitle", 14, DARK_BLUE, 0, 8),
        ("Heading 1", 17, NAVY, 18, 10),
        ("Heading 2", 13.5, BLUE, 14, 7),
        ("Heading 3", 11.5, DARK_BLUE, 10, 5),
    ]:
        s = doc.styles[name]
        s.font.name = FONT
        s._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor.from_string(color)
        s.font.bold = name != "Subtitle"
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    for name in ["List Bullet", "List Bullet 2", "List Number"]:
        s = doc.styles[name]
        s.font.name = FONT
        s._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        s.font.size = Pt(10.5)
        s.paragraph_format.space_after = Pt(4)
        s.paragraph_format.line_spacing = 1.25


def configure_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("PDMG 아키텍처 수립 가이드")
    set_run_font(r, size=8.5, color=GRAY)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("NSIGHT  |  ")
    set_run_font(r, size=8.5, color=GRAY)
    add_page_field(p)


def build():
    doc = Document()
    configure_styles(doc)
    configure_doc(doc)

    # Cover: editorial_cover pattern, restrained technical guide treatment.
    for _ in range(5):
        add_p(doc, "", after=12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_para(p, after=14)
    r = p.add_run("ARCHITECTURE FIELD GUIDE")
    set_run_font(r, size=11, bold=True, color=GOLD)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_para(p, after=10)
    r = p.add_run("PDMG 아키텍처 수립 가이드")
    set_run_font(r, size=30, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_para(p, after=28)
    r = p.add_run("비전에서 Runtime Evidence까지 연결하는 6단계 실무 방법론")
    set_run_font(r, size=14, color=DARK_BLUE)
    add_callout(doc, "핵심 원칙", "좋은 아키텍처는 잘 그린 그림이 아니라 책임·경계·실행규칙·실행 증거가 계속 정합되는 체계다.", "info")
    for _ in range(3):
        add_p(doc, "", after=12)
    add_p(doc, "기반 자료  |  NSIGHT PDMG 아키텍처 정의서 FINAL V5 (13개 장)", color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p(doc, "작성 기준일  |  2026-09-01", color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_page_break()
    add_heading(doc, "문서 정보", 1)
    add_table(doc, ["항목", "내용"], [
        ["문서명", "PDMG 아키텍처 수립 가이드"],
        ["목적", "아키텍처 수립·검토·승인·운영 검증을 하나의 반복 가능한 절차로 표준화"],
        ["주요 독자", "Enterprise/Solution/Application/Data/Infrastructure Architect, 개발·운영·보안·데이터 담당자"],
        ["적용 범위", "신규 구축, 현대화, 플랫폼 전환, 데이터·마케팅·BI·DR 아키텍처"],
        ["근거", "PDMG 아키텍처 정의서 FINAL V5 1~13장"],
        ["판정 체계", "PASS / CONDITIONAL / FAIL 및 Evidence 기반 Conformance"],
    ], [2200, 7160])
    add_heading(doc, "이 가이드의 사용법", 2)
    for x in [
        "프로젝트 착수 시 1부에서 공통 원칙과 대상 범위를 정렬한다.",
        "2부의 6단계를 순서대로 수행하되, 발견된 사실을 이전 단계로 환류한다.",
        "3부의 횡단 관심사와 도메인 적용 지침을 필요한 범위에 결합한다.",
        "각 단계 종료 시 Gate와 Evidence 패키지를 검토하고, CONDITIONAL 항목은 책임자와 기한을 명시한다.",
        "문서는 최종 산출물이 아니라 Source·Config·Deployment·Runtime 증거와 함께 갱신되는 기준선으로 관리한다.",
    ]:
        add_bullet(doc, x)

    doc.add_page_break()
    add_heading(doc, "목차", 1)
    add_toc_field(doc)

    add_section_intro(doc, 1, "아키텍처를 다시 정의하는 관점", "시스템 목록을 나열하는 방식에서 벗어나 책임, 경계, 실행 흐름, 증거를 연결하는 공통 언어를 확립한다.", "제1장 왜 다시 짓는가, 제2장 정보계 패러다임의 전환")
    add_heading(doc, "1. Inventory가 아니라 Architecture를 만든다", 1)
    add_p(doc, "Inventory는 서버·제품·모듈·테이블의 존재를 보여준다. Architecture는 각 요소가 왜 존재하며 어떤 책임을 지고, 어디에서 분리되고, 실패 시 어떤 영향을 주며, 실행 중 무엇으로 검증되는지를 설명한다.")
    add_table(doc, ["구분", "Inventory 관점", "Architecture 관점"], [
        ["관심사", "무엇이 있는가", "왜 있고 어떤 책임을 지는가"],
        ["경계", "프로젝트/서버/제품", "Runtime/Trust/Transaction/Failure 경계"],
        ["연결", "정적 의존 관계", "요청·데이터·배포·장애의 실제 흐름"],
        ["검증", "문서 존재 여부", "Source·Config·Deployment·Runtime Evidence"],
        ["변경", "문서 사후 수정", "Drift 탐지와 ADR을 통한 Closed Loop"],
    ], [1600, 3600, 4160])
    add_callout(doc, "판단 기준", "요소 이름을 지웠을 때 책임과 경계가 설명되지 않는 그림은 Architecture가 아니라 Inventory일 가능성이 높다.", "risk")

    add_heading(doc, "2. Responsibility-centric으로 전환한다", 1)
    add_p(doc, "모듈과 제품은 구현 단위일 뿐이다. 먼저 사용자 경험, 인증, 업무 실행, 데이터 접근, 외부 연계, 관측과 같은 책임을 정의하고, 책임 간 경계가 Runtime에서 어떻게 구현되는지 내려가야 한다.")
    for x in [
        "Application 이름보다 Capability와 책임을 먼저 정의한다.",
        "Source 모듈 경계와 Runtime 배포 경계를 별도로 그린다.",
        "FAST 응답과 DEEP 처리의 Workload 특성을 분리한다.",
        "Security와 Observability는 특정 박스가 아니라 전 구간의 횡단 관심사로 다룬다.",
        "ServiceId와 같은 일관된 식별자를 요청·정책·데이터·로그·배포의 Backbone으로 사용한다.",
    ]:
        add_bullet(doc, x)

    add_heading(doc, "3. Architecture Closed Loop", 1)
    add_table(doc, ["순환 단계", "핵심 활동", "확인 증거"], [
        ["Define", "원칙·책임·경계·규칙 정의", "Vision, Big Picture, Catalog, ADR"],
        ["Implement", "Logical을 Physical과 Mechanism으로 구현", "Source, Config, IaC, Pipeline"],
        ["Run", "실제 거래와 장애 시나리오 실행", "Trace, Metric, Log, Deployment record"],
        ["Compare", "정의와 실행의 차이 분석", "Conformance report, Drift list"],
        ["Decide", "GAP 수용·개선·예외 결정", "ADR, Action owner, Due date"],
        ["Update", "기준선과 자동 검증 규칙 갱신", "Versioned artifact, Machine-readable rule"],
    ], [1500, 3900, 3960])

    add_section_intro(doc, 2, "아키텍처 6단계 수립 방법론", "Vision에서 Runtime까지 내려가며 추상적 의도를 구현 가능한 구조와 검증 가능한 증거로 바꾼다.", "제3~9장")
    add_heading(doc, "6단계 전체 흐름", 1)
    add_table(doc, ["단계", "고정하는 것", "대표 산출물", "다음 단계 질문"], [
        ["1. Vision", "방향·원칙·품질 목표", "Vision statement, NFR, 제약", "무엇을 분리해야 하는가?"],
        ["2. Big Picture", "책임·경계·상호작용", "Context/Container view, Boundary map", "어떤 기술 역할이 필요한가?"],
        ["3. Logical", "기술 역할·상태·확장·실패 도메인", "Logical node catalog, Interface/Data contract", "어떤 자원에 배치할 것인가?"],
        ["4. Physical", "센터·네트워크·VM·JVM·DB 배치", "Deployment topology, Mapping matrix", "어떤 실행 규칙이 보장하는가?"],
        ["5. Mechanism", "필터·라우팅·스레드·트랜잭션·오류", "Mechanism spec, Policy matrix", "실제 거래에서 그대로 동작하는가?"],
        ["6. Runtime", "시간축·부하·장애·복구 행동", "Sequence/Trace, Evidence, Gate result", "정의와 실행의 차이는 무엇인가?"],
    ], [1150, 2350, 3000, 2860], font_size=8.7)
    add_callout(doc, "운영 원칙", "Top-down 설계와 Bottom-up 사실 확인을 반복한다. 단계는 직선형 문서 공정이 아니라 Evidence로 닫히는 순환 구조다.", "success")

    phases = [
        (1, "VISION", "기술 선택보다 먼저 비즈니스 방향, 품질 속성, 제약과 비타협 원칙을 고정한다.",
         [["비즈니스 목표", "성과·고객·운영 관점의 변화"], ["현행 GAP", "구조적 원인과 영향"], ["NFR", "가용성·성능·보안·복구 목표"], ["제약", "규제·예산·기술·일정·조직"]],
         {"questions":["왜 지금 아키텍처를 다시 정의해야 하는가?", "10년 동안 유지할 핵심 원칙은 무엇인가?", "FAST와 DEEP Workload의 목표는 어떻게 다른가?", "성공을 어떤 수치와 증거로 판정할 것인가?"],
          "steps":["이해관계자와 현행 문제를 증상이 아닌 구조적 원인으로 재정의한다.", "비즈니스 목표를 품질 속성과 Architecture Driver로 변환한다.", "원칙 간 충돌을 식별하고 우선순위를 결정한다.", "Scope, Assumption, Constraint, Out-of-scope를 명시한다.", "성공 지표와 Gate 기준을 정의하고 승인한다."],
          "decisions":["Application-centric에서 Responsibility-centric으로 전환할 범위", "FAST/DEEP 처리의 분리 수준", "가용성·RTO/RPO·보안·데이터 신선도 목표", "표준화와 예외 승인 원칙"]},
         [["Architecture Vision", "목표·원칙·범위·대상이 한 페이지에서 연결"], ["Quality Attribute Scenario", "자극·환경·대상·응답·측정값이 명시"], ["Constraint Register", "제약의 근거·영향·소유자 기록"], ["Initial ADR", "핵심 방향의 대안·선택·결과 기록"]],
         ("원칙과 측정 가능한 목표가 승인되고 충돌 우선순위가 정해짐", "핵심 원칙은 합의됐으나 일부 목표 수치·책임자가 미정", "목표가 제품명 나열에 머물거나 품질 목표가 측정 불가"),
         ("방향과 품질 목표가 먼저이고 기술은 그 결과로 선택된다.", "선호 제품을 먼저 정한 뒤 Vision을 사후 정당화한다."),
         ["승인된 Vision/NFR 문서", "현행 GAP와 목표의 추적표", "Architecture Driver별 ADR", "Stakeholder 승인 기록"]),
        (2, "BIG PICTURE", "사용자에서 데이터와 외부 시스템까지 책임과 경계를 한눈에 고정한다.",
         [["Vision/NFR", "드라이버와 품질 목표"], ["Actor/Channel", "사용자·채널·운영 주체"], ["Capability", "업무·플랫폼 책임"], ["Trust/Data boundary", "인증·인가·데이터 소유권"]],
         {"questions":["사용자 요청은 어떤 경계를 순서대로 통과하는가?", "인증과 인가는 어디에서 시작되고 어디까지 전파되는가?", "업무 실행과 데이터 접근 책임은 어떻게 분리되는가?", "외부 연계와 관측은 어느 경계에 속하는가?"],
          "steps":["Actor와 Channel을 식별하고 시스템 Context를 확정한다.", "UI Delivery, Authentication, Application Runtime, Data, External Integration 경계를 배치한다.", "경계별 책임·소유자·허용 인터페이스를 정의한다.", "Security와 Observability를 모든 흐름 위에 겹쳐 검토한다.", "Big Picture에서 보이는 GAP와 미결정을 ADR 후보로 등록한다."],
          "decisions":["Channel/UI Delivery 분리", "Authentication Boundary와 Trust 전파", "Application Runtime Boundary와 ServiceId 적용", "Data/External Integration 소유권"]},
         [["System Context", "Actor·시스템·외부 연계가 식별"], ["Big Picture", "책임·경계·주요 흐름이 단일 뷰에 표현"], ["Boundary Catalog", "각 경계의 책임·소유자·정책 기록"], ["GAP/ADR List", "누락·충돌·대안이 추적 가능"]],
         ("모든 핵심 흐름이 책임과 경계를 통과하며 소유자가 명확", "일부 외부 연계 또는 횡단 관심사가 임시 정의", "박스는 있으나 책임·Trust·Data 경계가 설명되지 않음"),
         ("책임과 경계를 먼저 정하고 제품·서버는 후속 단계에서 배치한다.", "현재 서버 구조를 그대로 Big Picture로 확대한다."),
         ["Big Picture 승인본", "Boundary별 RACI", "주요 흐름과 Trust/Data 표시", "미결정 ADR backlog"]),
        (3, "LOGICAL", "제품과 서버에서 독립적인 기술 역할, 상태, 확장 단위와 실패 도메인을 정의한다.",
         [["Big Picture", "경계와 주요 상호작용"], ["Capability catalog", "업무·플랫폼 책임"], ["Data contract", "소유·조회·변경 규칙"], ["NFR", "성능·가용성·보안 목표"]],
         {"questions":["각 Logical Node의 단일 책임은 무엇인가?", "상태는 어디에 존재하며 어떻게 일관성을 보장하는가?", "무엇을 독립적으로 확장하고 격리해야 하는가?", "실패가 어디까지 전파되고 어디서 차단되는가?"],
          "steps":["Capability를 Logical Node와 Service로 분해한다.", "각 Node의 책임·인터페이스·의존·데이터·소유자를 Catalog화한다.", "State model과 Transaction boundary를 정의한다.", "Scale unit과 Failure domain을 정의한다.", "Security boundary와 정책 적용 지점을 연결한다.", "Logical→Physical Handoff 조건을 명시한다."],
          "decisions":["Application Runtime 내부 역할 분리", "동기/비동기 및 FAST/DEEP 처리", "Stateful/Stateless 경계", "ServiceId와 Logical Node의 매핑"]},
         [["Logical Architecture", "제품 독립적인 역할·흐름 표현"], ["Logical Node Catalog", "책임·인터페이스·상태·확장·실패 도메인 기록"], ["Interface/Data Contract", "호출·데이터 소유·오류 규약 정의"], ["Traceability Matrix", "Vision→Boundary→Logical Node 추적"]],
         ("모든 Node가 단일 책임과 명시적 인터페이스·상태·실패 도메인을 가짐", "일부 Scale/Failure 가정이 성능 시험 전제", "제품명만 바뀐 물리 구조이거나 공유 상태·실패 전파가 불명확"),
         ("Logical은 기술 역할과 정책을 정의하고 Physical 선택과 독립적이다.", "VM/JVM/제품명을 Logical Node처럼 사용한다."),
         ["Logical Node Catalog", "Interface/Data contract", "State/Transaction model", "Scale/Failure domain 검토 기록"]),
        (4, "PHYSICAL", "Logical 역할을 센터·네트워크·VM·JVM·WAR·DB와 운영 가능한 배치 구조로 내린다.",
         [["Logical catalog", "역할·상태·확장·실패 도메인"], ["Infrastructure standard", "센터·네트워크·가상화·DB 기준"], ["Capacity estimate", "트래픽·동시성·데이터량"], ["DR target", "RTO/RPO·전환 범위"]],
         {"questions":["Logical Node는 어느 자원과 배포 단위에 매핑되는가?", "GSLB/L4/WEB/WAS 경로와 Trust 경계는 일치하는가?", "VM/JVM/WAR 분리가 Scale과 Failure Domain을 보장하는가?", "방화벽·포트·DB 경로가 실제로 실행 가능한가?"],
          "steps":["Logical→Physical Mapping matrix를 작성한다.", "센터·Zone·Network segment와 요청 경로를 배치한다.", "VM/JVM/WAR/Container 및 DB 인스턴스 분리 원칙을 적용한다.", "Port/Protocol/Firewall/Certificate 규칙을 정의한다.", "용량 후보와 확장 단위를 계산한다.", "HA/DR·백업·복구·운영 접근 경로를 검증한다.", "Physical traceability를 확정한다."],
          "decisions":["Main/DR 및 Local HA 토폴로지", "JVM/WAR Isolation 단위", "Data physical path와 복제", "Capacity headroom과 확장 방식"]},
         [["Deployment Topology", "센터부터 Runtime까지 배치·경로 표현"], ["Mapping Matrix", "Logical Node→자원→배포 단위 추적"], ["Network/Port Matrix", "방향·프로토콜·정책·소유자 기록"], ["Capacity/DR Design", "가정·수치·전환·복구 절차 정의"]],
         ("모든 Logical Node가 자원·경로·격리·용량·복구 구조에 추적", "일부 용량 수치는 부하시험 후 확정 조건", "공유 JVM/DB/Network가 의도한 실패·보안 경계를 무너뜨림"),
         ("물리 배치는 Logical 책임과 실패 도메인을 보존한다.", "현행 서버 수에 맞춰 Logical Node를 임의로 합친다."),
         ["CMDB/IaC/배포 정의", "Network·Firewall 승인", "Capacity 산정 근거", "HA/DR 테스트 계획"]),
        (5, "MECHANISM", "정적 구조를 실제로 움직이는 공통 실행규칙과 업무 확장 지점을 정의한다.",
         [["Logical/Physical", "역할·배치·경계"], ["Request contract", "ServiceId·Context·Identity"], ["Security policy", "인증·인가·Trust 규칙"], ["Runtime target", "Timeout·Thread·Transaction·Error 목표"]],
         {"questions":["요청 Context는 어디서 시작되고 어떻게 전달되는가?", "ServiceId는 어떻게 해석·라우팅·정책 적용되는가?", "Request Thread와 Worker Thread의 책임은 어떻게 분리되는가?", "Timeout·Transaction·Error·Logging의 순서는 무엇인가?"],
          "steps":["Framework 책임과 Business 확장 지점을 분리한다.", "DefaultFilter에서 Correlation/Context 생성 규칙을 정한다.", "SecurityFilterChain과 Trust Gate를 정의한다.", "ServiceId Resolution·TCF·Handler Registry를 연결한다.", "Worker/Queue/Timeout/Cancel 규칙을 정의한다.", "TransactionTemplate과 DB 접근 경계를 정한다.", "Error/Response/Logging 표준을 정의한다.", "정상·오류·과부하 시나리오로 메커니즘을 검토한다."],
          "decisions":["Framework vs Business 책임", "Thread/Queue/Timeout 모델", "Transaction 시작·종료·Rollback 경계", "503/504 및 표준 오류 응답 규칙"]},
         [["Mechanism Specification", "필터·라우팅·실행·트랜잭션·응답 순서 정의"], ["Policy Matrix", "ServiceId별 보안·Timeout·DB·Handler 정책"], ["Error Contract", "코드·HTTP 상태·로그·추적 규약"], ["Sequence Scenarios", "정상·Timeout·Overload·DB 오류 흐름"]],
         ("주요 거래가 동일 규칙으로 실행되고 오류·취소·Rollback까지 정의", "일부 비동기 취소/보상 규칙이 구현 검증 조건", "업무 코드가 공통 보안·Timeout·Transaction 규칙을 우회"),
         ("Framework는 공통 규칙을 강제하고 Business는 명시된 확장점만 사용한다.", "개별 업무가 Thread·Transaction·Security를 제각각 구현한다."),
         ["Framework source/config", "ServiceId policy registry", "오류·Timeout 테스트", "표준 로그/Trace schema"]),
        (6, "RUNTIME VALIDATION", "거래 한 건과 장애·과부하 시나리오를 시간축으로 펼쳐 설계와 실행의 일치를 증명한다.",
         [["Mechanism spec", "실행 순서와 정책"], ["Deployment", "실제 버전·자원·Config"], ["Test scenario", "정상·오류·부하·복구"], ["Evidence schema", "Trace·Metric·Log·판정 규칙"]],
         {"questions":["Request Thread와 Worker Thread가 의도대로 동작하는가?", "Timeout 504와 Overload 503이 구분되는가?", "DB Transaction과 취소·Rollback이 시간축에서 맞는가?", "Saturation이 Cascade로 확산되기 전에 차단되는가?"],
          "steps":["ServiceId별 대표 거래와 Correlation 기준을 선택한다.", "정상 요청을 End-to-End Trace로 재구성한다.", "Slow DB, Queue full, Worker saturation, 인증 실패, 센터 장애를 주입한다.", "Thread pool·Queue·Connection pool·GC·DB metric을 동시 관찰한다.", "예상 상태코드·Rollback·로그·알람과 실제 결과를 비교한다.", "Drift/GAP을 분류하고 ADR·개선 작업으로 환류한다.", "Gate를 판정하고 Evidence 패키지를 기준선으로 보관한다."],
          "decisions":["서비스별 SLO/SLI와 증거 보존", "503/504 전환 기준", "Saturation 보호와 Back-pressure", "Drift 수용·개선·예외 처리"]},
         [["Runtime Sequence/Trace", "시간축·Thread·ServiceId·DB·응답 연결"], ["Test Evidence", "조건·버전·Config·결과 재현 가능"], ["Conformance Report", "Expected/Actual/GAP/판정 기록"], ["ADR/Action Plan", "차이의 결정·소유자·완료 기준 명시"]],
         ("정상·오류·부하·복구 시나리오가 목표를 만족하고 재현 가능한 증거 보유", "핵심 시나리오는 통과했으나 일부 운영 자동화가 기한부 미완", "Trace 불가, 상태코드 혼재, Rollback 불명확 또는 Saturation cascade 발생"),
         ("실행 증거로 설계를 검증하고 차이를 다시 기준선에 반영한다.", "문서 승인만으로 완료 처리하거나 로그 양을 Evidence로 오인한다."),
         ["Trace/Metric/Log bundle", "배포 버전과 Config fingerprint", "실험 조건과 결과", "Conformance 판정·ADR·후속 조치"]),
    ]
    for phase in phases:
        add_phase(doc, *phase)

    workshop_sheets = [
        ("VISION 워크숍", "목표와 Architecture Driver 정렬", "Sponsor, Business Owner, Lead Architect, Security/Data/Operations Lead",
         ["현행의 구조적 문제가 무엇인지 합의", "비즈니스 목표를 품질 속성 시나리오로 변환", "원칙 충돌과 우선순위 결정", "Scope·Constraint·Out-of-scope 확정"],
         ["Vision one-pager", "Quality Attribute Scenario", "Constraint register", "초기 ADR"]),
        ("BIG PICTURE 워크숍", "책임·경계·주요 흐름 정렬", "Business/Domain Architect, Application, Security, Data, Integration Lead",
         ["Actor·Channel·External System 배치", "UI/Auth/App/Data/Integration Boundary 정의", "주요 요청·데이터·Trust 흐름 표시", "GAP와 ADR 후보 등록"],
         ["Big Picture", "Boundary catalog", "주요 흐름", "RACI/GAP list"]),
        ("LOGICAL 워크숍", "제품 독립적인 기술 역할 설계", "Solution/Application/Data Architect, Framework/Platform Lead",
         ["Capability를 Logical Node로 분해", "Node별 책임·인터페이스·상태 정의", "Scale Unit·Failure Domain 결정", "Security·Data Contract 연결"],
         ["Logical diagram", "Node catalog", "Contract", "Traceability matrix"]),
        ("PHYSICAL 워크숍", "운영 가능한 자원·경로·격리 설계", "Infrastructure/Network/DB/DR/Operations Architect, Application Lead",
         ["Logical→Physical 매핑", "센터·Network·GSLB/L4/WEB/WAS 경로", "VM/JVM/WAR/DB 격리", "Capacity·HA·DR·Port 검토"],
         ["Deployment topology", "Mapping matrix", "Network/port matrix", "Capacity/DR design"]),
        ("MECHANISM 워크숍", "공통 Runtime 실행규칙 설계", "Framework, Security, Application, DB, Observability Lead",
         ["Context·Security·ServiceId 흐름", "Handler·Worker·Queue·Timeout 결정", "Transaction·DB 경계", "Error·Response·Logging 표준"],
         ["Mechanism sequence", "Policy matrix", "Error contract", "Runtime test scenario"]),
        ("RUNTIME VALIDATION 워크숍", "정상·오류·부하·복구 Evidence 판정", "Architecture Board, QA/Performance, Operations, Security, Domain Owner",
         ["대표 ServiceId와 실험 조건 확정", "Trace·Metric·Log 동시 검토", "503/504·Rollback·Saturation 판정", "Drift·ADR·조치·Gate 확정"],
         ["Trace bundle", "Conformance report", "Gate result", "ADR/action plan"]),
    ]
    for title, objective, participants, agenda, deliverables in workshop_sheets:
        doc.add_page_break()
        add_heading(doc, title, 1)
        add_callout(doc, "워크숍 목표", objective, "info")
        add_heading(doc, "참여자", 2)
        add_p(doc, participants)
        add_heading(doc, "권장 진행안 (120분)", 2)
        agenda_rows = []
        times = ["0~15분", "15~45분", "45~90분", "90~120분"]
        for t, item in zip(times, agenda):
            agenda_rows.append([t, item, "결정/쟁점 기록"])
        add_table(doc, ["시간", "활동", "기록"], agenda_rows, [1500, 5360, 2500])
        add_heading(doc, "종료 시 확보할 산출물", 2)
        for item in deliverables:
            add_bullet(doc, "[ ] " + item)
        add_heading(doc, "Facilitator 확인", 2)
        for item in ["대안과 선택 이유가 기록됐는가?", "미결정 항목에 책임자와 기한이 있는가?", "다음 단계 입력과 Handoff 조건이 명확한가?", "Gate 판정에 필요한 Evidence가 합의됐는가?"]:
            add_bullet(doc, "[ ] " + item)
        add_callout(doc, "워크숍 산출 원칙", "회의록이 아니라 의사결정, 산출물 변경점, Evidence 요구사항을 남긴다.", "success")

    add_section_intro(doc, 3, "횡단 관심사와 도메인 적용", "6단계 방법론을 DR, 데이터, 마케팅, BI에 적용할 때 반드시 추가해야 할 설계 관점을 정리한다.", "제7장, 제10~12장")
    add_heading(doc, "1. DR 센터 활용 전략", 1)
    add_p(doc, "HA는 국지적 장애를 빠르게 흡수하는 구조이고, DR은 센터 단위 재난에서 서비스를 복구하는 구조다. 두 목표를 혼용하지 말고 장애 범위·RTO·RPO·전환·복귀를 각각 검증해야 한다.")
    add_table(doc, ["설계 항목", "필수 질문", "Evidence"], [
        ["Main Local HA", "프로세스·VM·장비 장애가 어느 범위에서 흡수되는가?", "HA test, health check, failover trace"],
        ["센터 전환", "Traffic·Application·Config·Key·DB가 어떤 순서로 전환되는가?", "Runbook, DNS/LB record, deployment trace"],
        ["DB 일관성", "복제 지연·손실·Split-brain을 어떻게 판정하는가?", "Replication metric, consistency check"],
        ["RTO/RPO", "서비스별 목표와 실제 측정값이 일치하는가?", "Timestamped recovery evidence"],
        ["Failback", "Main 복귀 시 데이터와 Traffic을 어떻게 재정렬하는가?", "Failback rehearsal, reconciliation result"],
    ], [1800, 4560, 3000], font_size=8.8)
    add_callout(doc, "DR PASS", "장비가 존재하는 상태가 아니라, 지정된 조건에서 사람이 실행 가능한 절차로 목표 시간과 데이터 손실 범위를 증명한 상태다.", "risk")

    add_heading(doc, "2. 데이터 플랫폼", 1)
    add_p(doc, "DB 제품이 아니라 데이터의 책임·소유권·흐름·품질·신선도·부하 격리를 설계한다. RDW는 FAST 업무의 일관성과 응답을 지키고, ADW는 DEEP 분석의 확장성과 활용을 극대화한다.")
    add_table(doc, ["관점", "RDW", "ADW/분석 영역"], [
        ["주요 목적", "실시간 업무 조회·변경", "분석·집계·탐색·학습"],
        ["우선 품질", "일관성·낮은 지연·트랜잭션", "처리량·확장성·신선도"],
        ["변경 규칙", "소유 서비스만 Direct DML", "계약된 Pipeline을 통한 적재"],
        ["연계", "ServiceId→Table 추적", "CDC와 ETL 목적 분리"],
        ["격리", "분석 부하로부터 보호", "업무 부하와 독립 확장"],
    ], [1900, 3730, 3730])
    for x in ["ServiceId→Handler→Repository→Table 추적성을 유지한다.", "Data Owner와 Direct DML 허용 주체를 명시한다.", "CDC는 변경 전달, ETL은 변환·품질·집계를 담당하도록 분리한다.", "Schema·품질·Lineage·Freshness를 Machine-readable contract로 관리한다."]:
        add_bullet(doc, x)

    add_heading(doc, "3. 마케팅 플랫폼", 1)
    add_p(doc, "마케팅 플랫폼은 배치 캠페인 도구가 아니라 고객 행동과 Context에 반응하는 Runtime Capability로 설계한다. Program은 비즈니스 목적을, ServiceId는 실행 가능한 서비스와 정책을 식별한다.")
    add_table(doc, ["영역", "설계 포인트", "검증 질문"], [
        ["Program/ServiceId", "캠페인 목적과 Runtime 서비스 연결", "모든 실행이 식별·추적되는가?"],
        ["Customer Context", "Identity·Consent·Segment·상태 결합", "신뢰 가능한 최신 Context인가?"],
        ["Business Runtime", "Rule·Decision·Action 분리", "업무 변경이 Framework를 우회하지 않는가?"],
        ["External Interface", "채널·파트너 호출 계약", "Timeout·재시도·중복·보상이 정의됐는가?"],
        ["Event/Kafka", "비동기 확장과 재처리", "순서·멱등·DLQ·Schema가 검증되는가?"],
    ], [1800, 4000, 3560])

    add_heading(doc, "4. BI 포털", 1)
    add_p(doc, "BI는 데이터 플랫폼 위에서 판단 속도를 높이는 소비 경계다. Report와 Self-BI를 구분하고, Data Contract·Security·Freshness SLA·Performance Isolation을 함께 설계한다.")
    add_table(doc, ["설계 항목", "완료 조건"], [
        ["BI Boundary", "업무 시스템과 분석 소비 책임이 분리됨"],
        ["RDW→ADW", "전달 방식·변환·신선도·재처리 규칙이 정의됨"],
        ["Data Contract", "Schema·정의·품질·소유·Lineage가 승인됨"],
        ["Report/Self-BI", "공식 지표와 탐색 영역의 책임·승인 경계가 명확"],
        ["BI Security", "사용자·역할·행/열·민감정보 정책이 적용됨"],
        ["Freshness SLA", "데이터셋별 목표와 측정·알람 기준이 존재"],
        ["Performance Isolation", "BI 부하가 RDW와 FAST 업무에 영향 없음"],
        ["BI Evidence", "조회·권한·신선도·성능 결과가 재현 가능"],
    ], [2500, 6860])

    add_section_intro(doc, 4, "표준화와 10년 지속 가능성", "문서·코드·설정·배포·실행의 정합성을 자동 검증하고 변화의 이유를 장기 보존한다.", "제13장 표준화와 10년 지속 가능성")
    add_heading(doc, "1. Naming Backbone", 1)
    add_p(doc, "동일한 식별자가 Architecture, Source, Config, Deployment, Runtime Evidence를 관통해야 한다. ServiceId, ApplicationId, ModuleId, Data Owner, Deployment Unit의 관계를 표준 사전으로 관리한다.")
    add_table(doc, ["식별자", "연결 대상", "관리 원칙"], [
        ["ServiceId", "요청·정책·Handler·DB·로그·Trace", "전 생명주기에서 불변·유일"],
        ["ApplicationId", "제품/업무 Capability·배포·소유 조직", "Capability와 책임을 명확히 표현"],
        ["Deployment Unit", "Artifact·환경·VM/JVM/Container", "Build Once / Promote 추적"],
        ["Data Owner", "Schema·Table·Contract·변경 권한", "Direct DML과 품질 책임 명시"],
    ], [1900, 4160, 3300])

    add_heading(doc, "2. Machine-readable Rule", 1)
    for x in [
        "Naming, Dependency, Security, Data access, Deployment, Observability 규칙을 정형 데이터로 표현한다.",
        "CI/CD와 정기 점검에서 규칙을 실행하고 위반은 Evidence와 함께 보고한다.",
        "예외는 암묵적으로 허용하지 않고 ADR 번호·소유자·만료일을 가진 정책으로 기록한다.",
        "문서의 표와 Catalog는 가능한 한 동일한 원천 데이터에서 생성한다.",
    ]:
        add_bullet(doc, x)

    add_heading(doc, "3. Build Once / Promote와 Config·Secret 분리", 1)
    add_table(doc, ["원칙", "정상 패턴", "금지 패턴"], [
        ["Build Once", "동일 Artifact를 환경 간 승격", "환경마다 재빌드"],
        ["Config", "환경별 값은 외부화·버전·검증", "Artifact 내부 하드코딩"],
        ["Secret", "Vault/KMS 등 보호 저장소와 회전", "소스·평문 Config 포함"],
        ["Deployment Trace", "Artifact hash·Config fingerprint·승인·시간 기록", "현재 버전을 운영자 기억에 의존"],
    ], [1700, 3830, 3830])

    add_heading(doc, "4. Observability에서 Evidence로", 1)
    add_p(doc, "로그·메트릭·트레이스를 많이 수집하는 것만으로는 Evidence가 되지 않는다. Architecture Decision과 Gate를 증명할 수 있도록 식별자, 시간, 버전, 설정, 기대 결과와 판정 규칙을 함께 보존해야 한다.")
    add_table(doc, ["Evidence 유형", "최소 포함 항목"], [
        ["Source", "Repository, commit, module, rule result"],
        ["Config", "Environment, version/fingerprint, effective value, owner"],
        ["Deployment", "Artifact hash, target, timestamp, approver, result"],
        ["Runtime", "ServiceId, correlation, trace, metric, log, response, DB result"],
        ["Decision", "Context, alternatives, choice, consequence, owner, review date"],
    ], [2100, 7260])

    add_heading(doc, "5. Drift와 ADR Closed Loop", 1)
    for x in [
        "정의된 기준선과 실제 Source·Config·Deployment·Runtime을 주기적으로 비교한다.",
        "차이를 오류, 계획된 변경, 승인된 예외, 측정 한계로 분류한다.",
        "구조적 결정은 ADR에 Context·대안·결정·결과·재검토 조건을 기록한다.",
        "GAP은 책임자·완료일·PASS 전환조건 없이 종료하지 않는다.",
    ]:
        add_number(doc, x)

    add_section_intro(doc, 5, "실무 체크리스트와 템플릿", "프로젝트에서 즉시 복사해 사용할 수 있는 검토 항목과 산출물 골격을 제공한다.", "13개 장의 Gate, Evidence, ADR, 정상/금지 패턴 통합")
    add_heading(doc, "1. 착수 체크리스트", 1)
    for x in [
        "[ ] 비즈니스 목표와 구조적 GAP가 연결되어 있다.", "[ ] 범위와 Out-of-scope가 승인되어 있다.", "[ ] 핵심 품질 속성이 측정 가능한 시나리오로 정의되어 있다.", "[ ] Architecture Owner와 도메인별 책임자가 지정되어 있다.", "[ ] 현행 Source·Config·Deployment·Runtime Evidence 접근 권한이 확보되어 있다.", "[ ] ADR 저장소와 Gate 운영 방식이 정해져 있다.",
    ]:
        add_bullet(doc, x)

    add_heading(doc, "2. 단계별 완료 체크리스트", 1)
    checklist_rows = []
    for phase, checks in [
        ("Vision", ["목표/NFR 측정 가능", "원칙 충돌 우선순위", "범위·제약·가정 승인"]),
        ("Big Picture", ["책임·경계·소유자", "주요 요청/데이터 흐름", "Security/Observability 횡단 검토"]),
        ("Logical", ["Node catalog", "State/Scale/Failure domain", "Interface/Data contract"]),
        ("Physical", ["Logical mapping", "Network/Port/Trust", "Capacity/HA/DR"]),
        ("Mechanism", ["Context/Security/ServiceId", "Thread/Timeout/Transaction", "Error/Response/Logging"]),
        ("Runtime", ["정상/오류/부하/복구", "재현 가능한 Evidence", "Conformance/ADR 환류"]),
    ]:
        checklist_rows.append([phase, " / ".join(f"[ ] {c}" for c in checks), "PASS / COND / FAIL"])
    add_table(doc, ["단계", "필수 확인", "판정"], checklist_rows, [1400, 6460, 1500], font_size=8.7)

    add_heading(doc, "3. Architecture Decision Record 템플릿", 1)
    add_table(doc, ["필드", "작성 내용"], [
        ["ADR ID / 제목", "유일 식별자와 결정의 핵심"], ["상태", "Proposed / Accepted / Superseded / Deprecated"],
        ["Context", "문제·제약·Architecture Driver"], ["Decision", "선택한 구조와 적용 범위"],
        ["Alternatives", "검토한 대안과 배제 이유"], ["Consequences", "긍정·부정 영향과 후속 과제"],
        ["Evidence", "결정을 검증할 Source/Config/Runtime 증거"], ["Review", "소유자·재검토 조건·일자"],
    ], [2400, 6960])

    add_heading(doc, "4. Logical Node Catalog 템플릿", 1)
    add_table(doc, ["필드", "설명"], [
        ["Node ID / 이름", "Naming Backbone을 따르는 유일 식별자"], ["책임", "한 문장으로 표현한 단일 책임"],
        ["Inbound/Outbound", "허용 인터페이스와 호출 계약"], ["State/Data", "보유·조회·변경 데이터와 소유권"],
        ["Scale Unit", "독립 확장 단위와 지표"], ["Failure Domain", "장애 범위·격리·복구"],
        ["Security", "Trust·인증·인가·민감정보"], ["Physical Mapping", "배포 자원과 Artifact"],
        ["Evidence", "Source·Config·Deployment·Runtime 추적 위치"],
    ], [2400, 6960])

    add_heading(doc, "5. Conformance 판정 템플릿", 1)
    add_table(doc, ["Rule/Decision", "Expected", "Actual Evidence", "GAP", "판정/조치"], [
        ["예: ServiceId별 Timeout", "정책 Registry 값 적용", "Trace + effective config", "일부 서비스 하드코딩", "COND / 제거 D+30"],
        ["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""],
    ], [1750, 1900, 2150, 1700, 1860], font_size=8.2)

    add_heading(doc, "6. 아키텍처 리뷰 회의 운영", 1)
    add_table(doc, ["순서", "내용", "시간 기준"], [
        ["1", "지난 Gate의 CONDITIONAL/GAP 조치 확인", "10분"],
        ["2", "이번 단계의 Decision과 대안 검토", "20분"],
        ["3", "Source·Config·Deployment·Runtime Evidence 확인", "20분"],
        ["4", "PASS/CONDITIONAL/FAIL 판정", "10분"],
        ["5", "ADR·책임자·기한·PASS 전환조건 확정", "10분"],
    ], [900, 6860, 1600])
    add_callout(doc, "회의 종료 조건", "모든 GAP에는 책임자, 완료일, Evidence, PASS 전환조건이 있어야 한다. '추후 협의'만 있는 항목은 종료된 결정이 아니다.", "risk")

    add_section_intro(doc, 6, "원문 13개 장 실무 참조 카드", "원문의 각 장을 실제 리뷰에서 빠르게 사용할 수 있도록 핵심 판단, 검토 질문, 정상·금지 패턴, Evidence 관점으로 압축한다.", "PDMG 아키텍처 정의서 FINAL V5 전체")
    chapter_cards = [
        ("1장. 왜 다시 짓는가", "Inventory를 넘어 책임·경계·실행·증거가 연결된 Architecture를 다시 세운다.",
         ["Inventory와 Architecture를 구분했는가?", "Module을 Responsibility로 재해석했는가?", "Source와 Runtime을 별도 View로 확인했는가?", "Timeout과 Transaction 경계가 충돌하지 않는가?", "Security GAP를 개별 기능이 아닌 구조 문제로 보았는가?", "ServiceId가 요청·정책·데이터·로그를 연결하는가?", "Logical에서 Physical까지 추적되는가?", "Closed Loop가 실제 운영 프로세스로 존재하는가?"],
         "ServiceId와 책임 중심의 추적 가능한 Architecture", "서버·제품 목록을 Architecture로 간주", ["Architecture map", "ServiceId trace", "Logical→Physical matrix", "Current GAP / ADR"]),
        ("2장. 정보계 패러다임의 전환", "Application·제품 중심에서 Responsibility·Capability·Runtime·Evidence 중심으로 전환한다.",
         ["Application-centric 한계가 구체적인 GAP로 정의됐는가?", "Responsibility-centric 구조가 합의됐는가?", "Module과 Runtime Boundary를 혼동하지 않는가?", "제품이 아니라 Capability를 정의했는가?", "FAST/DEEP Workload가 분리됐는가?", "Logical→Physical의 연결 규칙이 있는가?", "Logging을 판정 가능한 Evidence로 만들었는가?", "문서가 Closed Loop의 일부로 갱신되는가?"],
         "책임과 Workload에 맞춰 경계를 재편", "기존 Application 경계를 그대로 유지한 채 제품만 교체", ["Capability map", "Workload classification", "Runtime boundary", "Conformance evidence"]),
        ("3장. 아키텍처 6단계 수립 방법론", "Vision에서 Runtime까지 Top-down과 Bottom-up을 닫는 공통 수행 체계를 확립한다.",
         ["Vision이 방향과 품질 목표를 고정하는가?", "Big Picture가 책임과 경계를 고정하는가?", "Logical이 기술 역할을 정의하는가?", "Physical이 실제 자원으로 매핑되는가?", "Mechanism이 실행규칙을 정하는가?", "Runtime이 시간축으로 검증되는가?", "발견된 사실이 이전 단계에 환류되는가?", "각 단계가 Gate와 Evidence로 종료되는가?"],
         "6단계별 산출물·Gate·Evidence를 일관되게 운영", "단계별 문서만 만들고 상호 추적과 Runtime 검증을 생략", ["Six-stage roadmap", "Stage gate", "Traceability", "Runtime conformance"]),
        ("4장. Big Picture", "사용자에서 데이터·외부 연계까지 핵심 Boundary와 책임을 한 장에서 정렬한다.",
         ["User와 Channel Boundary가 분명한가?", "UI Delivery가 업무 실행과 분리되는가?", "Authentication Boundary와 Trust 전파가 보이는가?", "Application Runtime Boundary가 책임 단위로 구성됐는가?", "Data Boundary와 소유권이 정의됐는가?", "External Integration 책임이 격리됐는가?", "Security/Observability가 횡단 관심사로 표현됐는가?", "Big Picture에서 드러난 GAP가 등록됐는가?"],
         "책임·Trust·Data 경계가 흐름과 함께 표현", "현재 서버 배치를 확대해 그린 박스 다이어그램", ["Context/Container view", "Boundary catalog", "RACI", "GAP/ADR backlog"]),
        ("5장. 논리 아키텍처", "제품 독립적인 Logical Node의 책임·상태·확장·실패·보안 특성을 정의한다.",
         ["Application을 Capability로 분해했는가?", "Logical Node Catalog가 완성됐는가?", "Runtime Node 내부 책임이 분리됐는가?", "State Model과 소유권이 명확한가?", "Scale Unit이 독립적인가?", "Failure Domain이 제한되는가?", "Security Boundary가 정책 적용점과 연결되는가?", "Physical Handoff 조건이 정의됐는가?"],
         "각 Node가 단일 책임과 명시적 계약을 가짐", "제품명·VM명·JVM명을 Logical Node로 사용", ["Logical diagram", "Node catalog", "State/Scale/Failure model", "Interface contract"]),
        ("6장. 물리 아키텍처", "Logical 구조를 실제 센터·네트워크·VM·JVM·WAR·DB에 운영 가능하게 매핑한다.",
         ["Logical→Physical Mapping이 완전한가?", "GSLB/L4/WEB/WAS 경로가 명확한가?", "VM/JVM/WAR 분리 원칙이 있는가?", "JVM/WAR Isolation이 Failure Domain을 보존하는가?", "Data Physical Path가 추적되는가?", "Network/Firewall/Port가 승인됐는가?", "Capacity 후보가 가정과 함께 계산됐는가?", "Physical Traceability가 배포 정보와 연결되는가?"],
         "배치 구조가 Logical 책임과 격리를 보존", "자원 절감을 이유로 상이한 실패·보안 도메인을 무계획 통합", ["Deployment topology", "Mapping matrix", "Network/port matrix", "Capacity evidence"]),
        ("7장. DR 센터 활용 전략", "장비 보유가 아니라 목표 시간과 데이터 손실 범위를 증명하는 복구 구조를 만든다.",
         ["HA와 DR의 장애 범위가 구분되는가?", "Main Local HA가 자동·수동으로 검증됐는가?", "센터 장애 전환 순서가 명확한가?", "Application/Config/Key 복구가 재현 가능한가?", "DB Consistency 판정이 있는가?", "서비스별 RTO/RPO가 정의됐는가?", "Failover와 Failback을 모두 연습했는가?", "DR PASS가 Evidence로 판정되는가?"],
         "정기적인 전환·복귀 훈련과 Timestamped Evidence", "DR 장비 존재만으로 PASS 선언", ["DR runbook", "RTO/RPO measurement", "DB consistency result", "Failover/failback trace"]),
        ("8장. 메커니즘", "Framework 공통 규칙과 Business 확장점을 분리해 거래 실행의 일관성을 보장한다.",
         ["Framework와 Business 책임이 분리됐는가?", "DefaultFilter가 거래 Context를 시작하는가?", "SecurityFilterChain이 Trust Gate로 동작하는가?", "ServiceId Resolution과 TCF가 연결되는가?", "Handler Registry가 명시적으로 관리되는가?", "Worker/Timeout 규칙이 일관되는가?", "TransactionTemplate이 DB 경계를 보장하는가?", "Error/Response/Logging 표준이 적용되는가?"],
         "공통 실행규칙을 Framework가 강제", "업무별로 Security·Thread·Transaction·Error를 중복 구현", ["Mechanism sequence", "Policy registry", "Framework source/config", "Error/timeout tests"]),
        ("9장. 런타임 서비스", "정적인 구조를 거래 한 건의 시간축으로 펼쳐 정상·오류·과부하 행동을 검증한다.",
         ["Request Thread 역할이 제한되는가?", "Worker Thread가 격리·관리되는가?", "ServiceId Runtime Routing이 추적되는가?", "DB Runtime과 Transaction이 연결되는가?", "Timeout이 504로 일관되게 처리되는가?", "Overload가 503으로 조기에 차단되는가?", "Saturation Cascade 보호가 있는가?", "Runtime Evidence가 재현 가능한가?"],
         "Thread·Queue·DB·응답을 하나의 Trace로 연결", "지연과 과부하를 동일 오류로 처리해 장애를 확산", ["End-to-end trace", "Thread/queue metrics", "503/504 tests", "DB rollback evidence"]),
        ("10장. 데이터플랫폼", "DB 목록이 아니라 데이터 책임·소유·전달·품질·부하 격리를 설계한다.",
         ["Data Architecture와 DB 구성을 구분했는가?", "PDMG Current Data Access가 추적되는가?", "RDW 역할과 보호 기준이 명확한가?", "ADW 역할과 분석 확장이 명확한가?", "ServiceId→Table Trace가 가능한가?", "CDC와 ETL 목적이 분리됐는가?", "Data Ownership과 Direct DML 정책이 있는가?", "Data Evidence가 품질·신선도를 증명하는가?"],
         "RDW FAST와 ADW DEEP의 책임·부하 격리", "분석 편의를 위해 업무 DB에 무제한 직접 접근", ["ServiceId-table lineage", "Data contract", "CDC/ETL evidence", "Freshness/quality metrics"]),
        ("11장. 마케팅플랫폼", "Program 목적과 ServiceId 실행을 연결해 고객 행동에 반응하는 Runtime Platform을 만든다.",
         ["PDMG Runtime Reference를 따르는가?", "Program과 ServiceId 관계가 명확한가?", "Business Runtime 책임이 분리됐는가?", "Customer Context/Identity가 신뢰 가능한가?", "External Interface 계약이 있는가?", "Event/Kafka Target에 멱등·순서·재처리가 정의됐는가?", "MP↔mg Mapping이 추적되는가?", "Marketing Runtime Evidence가 존재하는가?"],
         "Context·Decision·Action이 식별·추적되는 이벤트 기반 실행", "캠페인 스크립트가 Identity·Consent·공통 Runtime을 우회", ["Program-ServiceId map", "Customer context contract", "Event schema/DLQ", "Marketing runtime trace"]),
        ("12장. BI 포탈", "신뢰 가능한 데이터 계약 위에서 공식 보고와 Self-BI를 안전하게 분리·운영한다.",
         ["BI Boundary가 업무 시스템과 분리되는가?", "RDW→ADW 흐름과 재처리가 정의됐는가?", "Data Contract가 승인됐는가?", "Report/Self-BI 책임이 구분되는가?", "BI Security가 행·열·민감정보까지 적용되는가?", "Freshness SLA가 데이터셋별로 측정되는가?", "Performance Isolation이 검증됐는가?", "BI Evidence가 조회·권한·신선도·성능을 증명하는가?"],
         "공식 지표·탐색 영역·보안·신선도의 명시적 계약", "BI 조회가 RDW 업무 부하와 데이터 소유권을 침해", ["BI boundary", "Data contract", "Security test", "Freshness/performance evidence"]),
        ("13장. 표준화와 10년 지속 가능성", "정의·구현·배포·실행의 정합성을 자동 검증하고 변화의 이유를 보존한다.",
         ["Naming Backbone이 전 생명주기를 연결하는가?", "핵심 규칙이 Machine-readable한가?", "Build Once / Promote가 지켜지는가?", "Config와 Secret이 분리되는가?", "Deployment Trace가 Artifact와 설정을 증명하는가?", "Observability가 Decision Evidence로 연결되는가?", "Drift와 ADR이 Closed Loop로 운영되는가?", "장기 전환 구간의 호환성과 제거 조건이 관리되는가?"],
         "규칙·증거·ADR이 자동 검증되는 지속적 Architecture", "문서와 실제 상태의 차이를 운영자의 기억으로 보정", ["Naming registry", "Rule execution report", "Deployment trace", "Drift/ADR register"]),
    ]
    for title, objective, questions, normal_pattern, forbidden_pattern, evidence_items in chapter_cards:
        doc.add_page_break()
        add_heading(doc, title, 1)
        add_callout(doc, "이 장의 목적", objective, "info")
        add_heading(doc, "리뷰 질문", 2)
        for q in questions:
            add_bullet(doc, "[ ] " + q)
        add_heading(doc, "정상 패턴과 금지 패턴", 2)
        add_table(doc, ["정상 패턴", "금지 패턴"], [[normal_pattern, forbidden_pattern]], [4680, 4680], header_fill=LIGHT_GRAY)
        add_heading(doc, "필수 Evidence", 2)
        add_table(doc, ["구분", "확보할 증거"], [[str(i + 1), item] for i, item in enumerate(evidence_items)], [900, 8460])
        add_callout(doc, "판정 메모", "Evidence가 없으면 PASS가 아니다. CONDITIONAL은 책임자·완료일·PASS 전환조건을 반드시 기록한다.", "risk")

    add_heading(doc, "부록 A. 핵심 용어", 1)
    glossary = [
        ["Architecture Backbone", "설계·구현·운영 증거를 관통하는 일관된 식별·추적 체계"],
        ["Boundary", "책임·Trust·Transaction·Failure·Data Ownership이 달라지는 지점"],
        ["Capability", "특정 제품에 종속되지 않은 비즈니스 또는 플랫폼 능력"],
        ["Closed Loop", "정의→구현→실행→비교→결정→갱신의 순환"],
        ["Conformance", "설계 기준과 실제 구현·실행이 일치하는지 Evidence로 판정하는 활동"],
        ["Drift", "승인된 Architecture 기준선과 실제 상태의 차이"],
        ["Evidence", "조건·버전·식별자·결과가 있어 재현·판정 가능한 증거"],
        ["FAST/DEEP", "낮은 지연의 실시간 처리와 처리량·분석 중심의 심층 처리 구분"],
        ["Failure Domain", "하나의 장애가 함께 영향을 미칠 수 있는 최대 범위"],
        ["Logical Node", "제품·서버와 독립적으로 정의된 기술적 역할 단위"],
        ["Mechanism", "정적 구조를 실제로 움직이는 필터·라우팅·스레드·트랜잭션·오류 규칙"],
        ["ServiceId", "서비스 실행과 정책·데이터·로그·Trace를 연결하는 식별자"],
    ]
    add_table(doc, ["용어", "정의"], glossary, [2500, 6860])

    add_heading(doc, "부록 B. 근거 자료 매핑", 1)
    source_rows = [
        ["1장", "왜 다시 짓는가", "Inventory vs Architecture, ServiceId Backbone, Closed Loop"],
        ["2장", "정보계 패러다임의 전환", "Responsibility/Capability/Runtime/Evidence 중심 전환"],
        ["3장", "아키텍처 6단계 수립 방법론", "Vision→Big Picture→Logical→Physical→Mechanism→Runtime"],
        ["4장", "Big Picture", "User/Channel/UI/Auth/App/Data/External Boundary"],
        ["5장", "논리 아키텍처", "Logical Node, State, Scale, Failure, Security"],
        ["6장", "물리 아키텍처", "Mapping, Traffic path, Isolation, Capacity, Traceability"],
        ["7장", "DR 센터 활용 전략", "HA/DR, RTO/RPO, Failover/Failback"],
        ["8장", "메커니즘", "Filter, Security, ServiceId, Worker, Transaction, Error"],
        ["9장", "런타임 서비스", "Thread, Timeout 504, Overload 503, Saturation, Evidence"],
        ["10장", "데이터플랫폼", "RDW/ADW, CDC/ETL, Ownership, Data Evidence"],
        ["11장", "마케팅플랫폼", "Program/ServiceId, Customer Context, Event/Kafka"],
        ["12장", "BI 포탈", "Data Contract, Security, Freshness, Isolation"],
        ["13장", "표준화와 10년 지속 가능성", "Naming, Rules, Promote, Trace, Drift, ADR"],
    ]
    add_table(doc, ["근거", "주제", "가이드 반영 내용"], source_rows, [900, 2800, 5660], font_size=8.4)
    add_p(doc, "본 가이드는 상기 자료의 핵심 개념과 검증 체계를 실무 수행 순서에 맞게 재구성한 문서다.", italic=True, color=GRAY)

    # Core properties
    doc.core_properties.title = "PDMG 아키텍처 수립 가이드"
    doc.core_properties.subject = "PDMG 아키텍처 6단계 수립 방법론 및 실무 체크리스트"
    doc.core_properties.author = "NSIGHT"
    doc.core_properties.keywords = "PDMG, Architecture, Vision, Big Picture, Logical, Physical, Mechanism, Runtime, Evidence"
    doc.core_properties.comments = "Based on NSIGHT PDMG Architecture Book FINAL V5 chapters 1-13"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
